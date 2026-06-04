"""
excel_builder.py
----------------
Builds Excel files from extracted financial data dicts.
Works for any company — generic structure.

Two outputs:
  1. Per-period files: {period}_financials.xlsx  (Raw_Extract, Structured_Data, Metadata)
  2. Forecast model:  forecast_model.xlsx        (Historical_Model, Drivers, Forecast,
                                                  Scenarios, Validation, Commentary, Sources,
                                                  Raw_Extract)
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ── Color constants ────────────────────────────────────────────────────────────
BLUE_INPUT  = "FF0000FF"
BLACK_FORM  = "FF000000"
WHITE_TEXT  = "FFFFFFFF"
HEADER_BG   = "FF1F4E79"
ALT_ROW_BG  = "FFF2F2F2"

FMT_CURRENCY = '#,##0;(#,##0);"-"'
FMT_PCT      = "0.0%;(0.0%);-"
FMT_2DP      = "0.00;(0.00);-"

# ── Helper utilities ───────────────────────────────────────────────────────────
def _font(bold=False, color=BLACK_FORM, size=10):
    return Font(name="Arial", bold=bold, color=color, size=size)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _header_row(ws, row, values, bg=HEADER_BG, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
        c.fill = _fill(bg)
        c.alignment = Alignment(horizontal="center", wrap_text=True)

def _set_col_widths(ws, widths):
    """widths: dict of {col_letter: width}"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def _safe(val, default=0):
    return val if val is not None else default


# ── Per-period Excel file ──────────────────────────────────────────────────────
def build_period_file(data: dict) -> bytes:
    """
    Build a single-period Excel file in memory and return as bytes.
    data: normalized dict from extractor.normalize_extracted_data()
    """
    wb = Workbook()

    # ── Raw_Extract ─────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Raw_Extract"
    _set_col_widths(ws, {"A": 50, "B": 15, "C": 20, "D": 40, "E": 30, "F": 18, "G": 10, "H": 12, "I": 12})

    raw_headers = ["source_file", "period", "statement_type", "original_label",
                   "normalized_label", "value", "unit", "page_number", "confidence"]
    _header_row(ws, 1, raw_headers)

    sf   = data.get("source_file", "")
    per  = data.get("period", "")
    unit = data.get("unit", "$")

    raw_rows = [
        # Income Statement
        (sf, per, "Income Statement", "Revenue", "Revenue", data.get("revenue"), unit, "-", "High"),
        (sf, per, "Income Statement", "Cost of goods sold / COGS", "COGS", data.get("cogs"), unit, "-", "High"),
        (sf, per, "Income Statement", "Gross profit / margin", "Gross Profit", data.get("gross_profit"), unit, "-", "High"),
        (sf, per, "Income Statement", "Operating expenses", "Operating Expenses", data.get("opex_operating"), unit, "-", "High"),
        (sf, per, "Income Statement", "Depreciation & amortization", "D&A", data.get("opex_amortization"), unit, "-", "High"),
        (sf, per, "Income Statement", "Rent / occupancy", "Rent & Occupancy", data.get("opex_rent"), unit, "-", "High"),
        (sf, per, "Income Statement", "Total operating expenses", "Total OPEX", data.get("total_opex"), unit, "-", "High"),
        (sf, per, "Income Statement", "Operating income / EBIT", "EBIT", data.get("operating_income"), unit, "-", "High"),
        (sf, per, "Income Statement", "Government assistance", "Government Assistance", data.get("government_assistance"), unit, "-", "High"),
        (sf, per, "Income Statement", "Income before taxes", "Pre-Tax Income", data.get("income_before_tax"), unit, "-", "High"),
        (sf, per, "Income Statement", "Current income tax", "Current Tax", data.get("current_tax"), unit, "-", "High"),
        (sf, per, "Income Statement", "Deferred income tax", "Deferred Tax", data.get("future_tax"), unit, "-", "High"),
        (sf, per, "Income Statement", "Net income", "Net Income", data.get("net_income"), unit, "-", "High"),
        # Cash Flow
        (sf, per, "Cash Flow", "Operating cash flow", "Operating Cash Flow", data.get("ocf"), unit, "-", "High"),
        (sf, per, "Cash Flow", "Capital expenditures / purchase of assets", "Capex", data.get("capex"), unit, "-", "High"),
        (sf, per, "Cash Flow", "Depreciation & amortization (add-back)", "D&A (non-cash)", data.get("amortization"), unit, "-", "High"),
        (sf, per, "Cash Flow", "Investing cash flow", "Investing CF", data.get("investing_cf"), unit, "-", "High"),
        (sf, per, "Cash Flow", "Financing cash flow", "Financing CF", data.get("financing_cf"), unit, "-", "High"),
        # Balance Sheet
        (sf, per, "Balance Sheet", "Cash and cash equivalents", "Cash", data.get("cash"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Accounts receivable", "Accounts Receivable", data.get("accounts_receivable"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Inventories", "Inventory", data.get("inventories"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "PP&E / capital assets (net)", "PP&E Net", data.get("capital_assets"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Total assets", "Total Assets", data.get("total_assets"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Accounts payable", "Accounts Payable", data.get("accounts_payable"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Total debt / long-term debt", "Total Debt", data.get("total_debt"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Retained earnings", "Retained Earnings", data.get("retained_earnings"), unit, "-", "High"),
        (sf, per, "Balance Sheet", "Total equity", "Total Equity", data.get("total_equity"), unit, "-", "High"),
    ]

    for i, row in enumerate(raw_rows, start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = _font(color=BLUE_INPUT if isinstance(v, (int, float)) and v is not None else BLACK_FORM)
        ws.cell(row=i, column=6).number_format = FMT_CURRENCY

    # ── Structured_Data ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Structured_Data")
    _set_col_widths(ws2, {"A": 38, "B": 22})

    company = data.get("company_name", "Company")
    period_lbl = data.get("reporting_period_label", per)
    ws2.cell(row=1, column=1, value=f"{company} — {period_lbl}").font = _font(bold=True, size=12)
    ws2.cell(row=2, column=2, value=f"Value ({unit})").font = _font(bold=True)

    def write_section(title, items, start_row):
        r = start_row
        # Section header
        c = ws2.cell(row=r, column=1, value=title)
        c.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
        c.fill = _fill(HEADER_BG)
        ws2.cell(row=r, column=2).fill = _fill(HEADER_BG)
        r += 1
        row_refs = {}
        for label, val, fmt, key in items:
            if label == "":
                r += 1
                continue
            ws2.cell(row=r, column=1, value=label).font = _font()
            vc = ws2.cell(row=r, column=2)
            if callable(val):
                fval = val(row_refs)
                vc.value = fval
                vc.font = _font(color=BLACK_FORM)
            elif val is None:
                vc.value = None
            else:
                vc.value = val
                vc.font = _font(color=BLUE_INPUT)
            vc.number_format = fmt
            if key:
                row_refs[key] = r
            r += 1
        return r, row_refs

    # Determine gross margin formula row dynamically
    # We'll write items sequentially and track rows
    r = 3
    # Income Statement
    inc_title = "INCOME STATEMENT"
    c = ws2.cell(row=r, column=1, value=inc_title)
    c.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
    c.fill = _fill(HEADER_BG)
    ws2.cell(row=r, column=2).fill = _fill(HEADER_BG)
    r += 1

    row_refs = {}

    def write_line(label, val, fmt, key=None):
        nonlocal r
        if label == "":
            r += 1
            return
        ws2.cell(row=r, column=1, value=label).font = _font()
        vc = ws2.cell(row=r, column=2)
        vc.value = val
        if isinstance(val, str) and val.startswith("="):
            vc.font = _font(color=BLACK_FORM)
        else:
            vc.font = _font(color=BLUE_INPUT if val is not None else BLACK_FORM)
        vc.number_format = fmt
        if key:
            row_refs[key] = r
        r += 1

    rev_r = r
    write_line("Revenue", data.get("revenue"), FMT_CURRENCY, "revenue")
    cogs_r = r
    write_line("Cost of Goods Sold (COGS)", data.get("cogs"), FMT_CURRENCY, "cogs")
    gp_r = r
    write_line("Gross Profit", f"=B{rev_r}-B{cogs_r}", FMT_CURRENCY, "gp")
    write_line("Gross Margin %", f"=B{gp_r}/B{rev_r}", FMT_PCT, "gm_pct")
    write_line("", None, "")
    opex_r = r
    write_line("Operating Expenses (excl. D&A)", data.get("opex_operating"), FMT_CURRENCY, "opex")
    da_r = r
    write_line("Depreciation & Amortization", data.get("opex_amortization"), FMT_CURRENCY, "da")
    rent_r = r
    write_line("Rent & Occupancy", data.get("opex_rent"), FMT_CURRENCY, "rent")
    totopex_r = r
    write_line("Total OPEX", f"=B{opex_r}+B{da_r}+B{rent_r}", FMT_CURRENCY, "totopex")
    write_line("", None, "")
    ebit_r = r
    write_line("Operating Income (EBIT)", f"=B{gp_r}-B{totopex_r}", FMT_CURRENCY, "ebit")
    write_line("EBIT Margin %", f"=B{ebit_r}/B{rev_r}", FMT_PCT)
    write_line("Government Assistance", data.get("government_assistance"), FMT_CURRENCY)
    write_line("Income Before Tax", data.get("income_before_tax"), FMT_CURRENCY)
    write_line("Current Tax", data.get("current_tax"), FMT_CURRENCY)
    write_line("Deferred Tax", data.get("future_tax"), FMT_CURRENCY)
    ni_r = r
    write_line("Net Income", data.get("net_income"), FMT_CURRENCY, "ni")
    write_line("Net Margin %", f"=B{ni_r}/B{rev_r}", FMT_PCT)
    write_line("", None, "")

    # Cash Flow section
    cf_title = ws2.cell(row=r, column=1, value="CASH FLOW STATEMENT")
    cf_title.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
    cf_title.fill = _fill(HEADER_BG)
    ws2.cell(row=r, column=2).fill = _fill(HEADER_BG)
    r += 1

    ocf_r = r
    write_line("Operating Cash Flow", data.get("ocf"), FMT_CURRENCY, "ocf")
    capex_r = r
    write_line("Capital Expenditures (Capex)", data.get("capex"), FMT_CURRENCY, "capex")
    write_line("Free Cash Flow", f"=B{ocf_r}-B{capex_r}", FMT_CURRENCY)
    write_line("D&A (non-cash add-back)", data.get("amortization"), FMT_CURRENCY)
    write_line("Investing Cash Flow", data.get("investing_cf"), FMT_CURRENCY)
    write_line("Financing Cash Flow", data.get("financing_cf"), FMT_CURRENCY)
    write_line("", None, "")

    # Balance Sheet section
    bs_title = ws2.cell(row=r, column=1, value="BALANCE SHEET")
    bs_title.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
    bs_title.fill = _fill(HEADER_BG)
    ws2.cell(row=r, column=2).fill = _fill(HEADER_BG)
    r += 1

    write_line("Cash", data.get("cash"), FMT_CURRENCY)
    write_line("Accounts Receivable", data.get("accounts_receivable"), FMT_CURRENCY)
    write_line("Inventories", data.get("inventories"), FMT_CURRENCY)
    write_line("PP&E / Capital Assets (Net)", data.get("capital_assets"), FMT_CURRENCY)
    write_line("Total Assets", data.get("total_assets"), FMT_CURRENCY)
    write_line("", None, "")
    write_line("Accounts Payable", data.get("accounts_payable"), FMT_CURRENCY)
    write_line("Total Debt", data.get("total_debt"), FMT_CURRENCY)
    write_line("Total Liabilities", data.get("total_liabilities"), FMT_CURRENCY)
    write_line("", None, "")
    write_line("Retained Earnings", data.get("retained_earnings"), FMT_CURRENCY)
    write_line("Total Equity", data.get("total_equity"), FMT_CURRENCY)

    # ── Metadata ─────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Metadata")
    _set_col_widths(ws3, {"A": 30, "B": 60})
    _header_row(ws3, 1, ["Field", "Value"], start_col=1)

    meta = [
        ("Extraction Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Source File", data.get("source_file", "")),
        ("Company", data.get("company_name", "")),
        ("Reporting Period", data.get("reporting_period_label", "")),
        ("Period Code", data.get("period", "")),
        ("Currency", data.get("currency", "")),
        ("Unit", data.get("unit", "")),
        ("Audit Status", data.get("audit_status", "")),
        ("Notes", data.get("notes", "")),
        ("Extraction Method", "Claude API (Anthropic)"),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        ws3.cell(row=i, column=1, value=k).font = _font(bold=True)
        ws3.cell(row=i, column=2, value=v).font = _font()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Forecast Model ─────────────────────────────────────────────────────────────
def build_forecast_model(periods_data: list[dict]) -> bytes:
    """
    Build a multi-period forecast model from a list of normalized period dicts.
    periods_data: sorted list of dicts (oldest → newest)
    Returns Excel file as bytes.
    """
    wb = Workbook()
    periods_data = sorted(periods_data, key=lambda d: d.get("period", ""))
    n = len(periods_data)
    period_labels = [d.get("period", f"Period {i+1}") for i, d in enumerate(periods_data)]
    company = periods_data[0].get("company_name", "Company") if periods_data else "Company"

    def col(i):
        return get_column_letter(i + 2)  # data starts at column B (index 2)

    # ── Historical_Model ────────────────────────────────────────────────────
    ws_h = wb.active
    ws_h.title = "Historical_Model"
    _set_col_widths(ws_h, {"A": 38, **{get_column_letter(i+2): 18 for i in range(n)}})

    ws_h.cell(row=1, column=1, value=f"{company} — Historical Financial Model").font = _font(bold=True, size=12)
    ws_h.cell(row=2, column=1, value="Blue = source values  |  Black = formulas").font = _font()

    _header_row(ws_h, 3, ["Line Item"] + period_labels)

    # Track row numbers for formula references
    rmap = {}
    r = 4

    def h_section(title):
        nonlocal r
        for col_i in range(n + 1):
            c = ws_h.cell(row=r, column=col_i + 1, value=title if col_i == 0 else None)
            c.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
            c.fill = _fill(HEADER_BG)
        r += 1

    def h_input(label, key, fmt=FMT_CURRENCY):
        nonlocal r
        ws_h.cell(row=r, column=1, value=label).font = _font()
        rmap[key] = r
        for i, d in enumerate(periods_data):
            val = d.get(key)
            c = ws_h.cell(row=r, column=i + 2, value=val)
            c.font = _font(color=BLUE_INPUT)
            c.number_format = fmt
        r += 1

    def h_formula(label, formula_template, key=None, fmt=FMT_CURRENCY):
        """formula_template uses {col} placeholder for the column letter."""
        nonlocal r
        ws_h.cell(row=r, column=1, value=label).font = _font()
        if key:
            rmap[key] = r
        for i in range(n):
            cl = get_column_letter(i + 2)
            f = formula_template.replace("{col}", cl)
            c = ws_h.cell(row=r, column=i + 2, value=f)
            c.font = _font(color=BLACK_FORM)
            c.number_format = fmt
        r += 1

    def h_blank():
        nonlocal r
        r += 1

    h_section("── INCOME STATEMENT ──")
    h_input("Revenue", "revenue")
    h_input("Cost of Goods Sold", "cogs")
    h_formula("Gross Profit", "={col}" + str(0) + "-{col}" + str(0),  # placeholder, fixed below
              key="gp")

    # Fix gross profit formula after we know rev_r and cogs_r
    rev_r  = rmap["revenue"]
    cogs_r = rmap["cogs"]
    gp_r   = rmap["gp"]
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=gp_r, column=i+2).value = f"={cl}{rev_r}-{cl}{cogs_r}"

    h_formula("Gross Margin %", "={col}" + str(gp_r) + "/{col}" + str(rev_r),
              key="gm_pct", fmt=FMT_PCT)
    # fix gm formula
    gm_r = rmap["gm_pct"]
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=gm_r, column=i+2).value = f"={cl}{gp_r}/{cl}{rev_r}"

    h_blank()
    h_input("Operating Expenses (excl. D&A)", "opex_operating")
    h_input("Depreciation & Amortization", "opex_amortization")
    h_input("Rent & Occupancy", "opex_rent")

    opex_r = rmap["opex_operating"]
    da_r   = rmap["opex_amortization"]
    rent_r = rmap["opex_rent"]
    totopex_r = r  # will be next row
    h_formula("Total OPEX", "{col}0", key="totopex")
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=totopex_r, column=i+2).value = f"={cl}{opex_r}+{cl}{da_r}+{cl}{rent_r}"
        ws_h.cell(row=totopex_r, column=i+2).number_format = FMT_CURRENCY

    h_blank()
    ebit_r = r
    h_formula("Operating Income (EBIT)", "{col}0", key="ebit")
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=ebit_r, column=i+2).value = f"={cl}{gp_r}-{cl}{totopex_r}"

    ebit_pct_r = r
    h_formula("EBIT Margin %", "{col}0", fmt=FMT_PCT)
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=ebit_pct_r, column=i+2).value = f"={cl}{ebit_r}/{cl}{rev_r}"
        ws_h.cell(row=ebit_pct_r, column=i+2).number_format = FMT_PCT

    h_blank()
    h_input("Government Assistance", "government_assistance")
    ni_r = r
    h_input("Net Income", "net_income")
    ni_pct_r = r
    h_formula("Net Margin %", "{col}0", fmt=FMT_PCT)
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=ni_pct_r, column=i+2).value = f"={cl}{ni_r}/{cl}{rev_r}"
        ws_h.cell(row=ni_pct_r, column=i+2).number_format = FMT_PCT

    h_blank()
    h_section("── CASH FLOW ──")
    ocf_r = r
    h_input("Operating Cash Flow", "ocf")
    capex_r = r
    h_input("Capital Expenditures", "capex")
    fcf_r = r
    h_formula("Free Cash Flow", "{col}0", key="fcf")
    for i in range(n):
        cl = get_column_letter(i + 2)
        ws_h.cell(row=fcf_r, column=i+2).value = f"={cl}{ocf_r}-{cl}{capex_r}"
    h_input("D&A", "amortization")

    h_blank()
    h_section("── BALANCE SHEET ──")
    h_input("Cash", "cash")
    h_input("Accounts Receivable", "accounts_receivable")
    h_input("Inventories", "inventories")
    h_input("PP&E (Net)", "capital_assets")
    h_input("Total Assets", "total_assets")
    h_blank()
    h_input("Accounts Payable", "accounts_payable")
    h_input("Total Debt", "total_debt")
    h_input("Total Equity", "total_equity")

    # ── Drivers ─────────────────────────────────────────────────────────────
    ws_d = wb.create_sheet("Drivers")
    _set_col_widths(ws_d, {"A": 35, **{get_column_letter(i+2): 16 for i in range(n+2)}})
    ws_d.cell(row=1, column=1, value=f"{company} — Key Financial Drivers").font = _font(bold=True, size=12)

    drv_headers = ["Driver"] + period_labels + ["YoY (Latest)"]
    _header_row(ws_d, 2, drv_headers)

    drv_items = [
        ("Revenue Growth %",      "revenue",        "growth"),
        ("COGS % of Revenue",     "cogs",           "pct_rev"),
        ("Gross Margin %",        "gp",             "pct_rev"),
        ("OPEX % of Revenue",     "totopex",        "pct_rev"),
        ("EBIT Margin %",         "ebit",           "pct_rev"),
        ("Capex % of Revenue",    "capex",          "pct_rev"),
        ("OCF Conversion (OCF/NI)", "ocf",          "pct_ni"),
    ]

    for dr, (drv_label, key, calc_type) in enumerate(drv_items, start=3):
        ws_d.cell(row=dr, column=1, value=drv_label).font = _font()
        src_row = rmap.get(key)
        rev_row = rmap.get("revenue")
        ni_row  = rmap.get("net_income", ni_r)
        for i in range(n):
            cl = get_column_letter(i + 2)
            vc = ws_d.cell(row=dr, column=i + 2)
            if src_row and rev_row:
                if calc_type == "growth" and i == 0:
                    vc.value = None
                elif calc_type == "growth" and i > 0:
                    prev_cl = get_column_letter(i + 1)
                    vc.value = f"=Historical_Model!{cl}{src_row}/Historical_Model!{prev_cl}{src_row}-1"
                    vc.font = _font(color=BLACK_FORM)
                elif calc_type == "pct_rev":
                    vc.value = f"=Historical_Model!{cl}{src_row}/Historical_Model!{cl}{rev_row}"
                    vc.font = _font(color=BLACK_FORM)
                elif calc_type == "pct_ni":
                    vc.value = f"=Historical_Model!{cl}{src_row}/Historical_Model!{cl}{ni_row}"
                    vc.font = _font(color=BLACK_FORM)
            vc.number_format = FMT_PCT
        # YoY latest (last col - prev col)
        if n >= 2:
            last_cl = get_column_letter(n + 1)
            prev_cl = get_column_letter(n)
            yoy_c = ws_d.cell(row=dr, column=n + 2)
            yoy_c.value = f"={last_cl}{dr}-{prev_cl}{dr}"
            yoy_c.font = _font(color=BLACK_FORM)
            yoy_c.number_format = FMT_PCT

    # ── Forecast ─────────────────────────────────────────────────────────────
    ws_f = wb.create_sheet("Forecast")
    fc_years = 3
    _set_col_widths(ws_f, {"A": 38, **{get_column_letter(i+2): 18 for i in range(fc_years)}})

    ws_f.cell(row=1, column=1, value=f"{company} — Forecast Model").font = _font(bold=True, size=12)
    ws_f.cell(row=2, column=1, value="Change assumptions (rows below) to update all outputs.").font = _font()

    # Derive base year label
    base_period = period_labels[-1] if period_labels else "2020"
    try:
        base_year = int(base_period[:4])
    except ValueError:
        base_year = 2020
    fc_labels = [f"{base_year + i + 1}E" for i in range(fc_years)]
    _header_row(ws_f, 3, ["Line Item"] + fc_labels)

    # Assumption rows
    asmp_start = 28
    ws_f.cell(row=asmp_start, column=1, value="── FORECAST ASSUMPTIONS (edit these) ──").font = Font(name="Arial", bold=True, color=WHITE_TEXT)
    ws_f.cell(row=asmp_start, column=1).fill = _fill(HEADER_BG)
    for col_i in range(2, fc_years + 2):
        ws_f.cell(row=asmp_start, column=col_i).fill = _fill(HEADER_BG)

    asmp_labels = ["Revenue Growth Rate", "COGS % of Revenue", "OPEX % of Revenue", "Capex % of Revenue"]
    # Compute defaults from last year
    last = periods_data[-1] if periods_data else {}
    rev_last = _safe(last.get("revenue"), 1)

    def avg_pct(key, denom_key="revenue"):
        vals = [_safe(d.get(key)) / max(_safe(d.get(denom_key)), 1) for d in periods_data if d.get(key) and d.get(denom_key)]
        return sum(vals) / len(vals) if vals else 0

    defaults = [0.05, avg_pct("cogs"), avg_pct("total_opex"), avg_pct("capex")]
    for ar, (lbl, default) in enumerate(zip(asmp_labels, defaults)):
        ws_f.cell(row=asmp_start + 1 + ar, column=1, value=lbl).font = _font()
        for col_i in range(2, fc_years + 2):
            c = ws_f.cell(row=asmp_start + 1 + ar, column=col_i, value=round(default, 4))
            c.font = _font(color=BLUE_INPUT)
            c.number_format = FMT_PCT

    rev_asmp_r  = asmp_start + 1
    cogs_asmp_r = asmp_start + 2
    opex_asmp_r = asmp_start + 3
    cap_asmp_r  = asmp_start + 4

    # Forecast rows
    fc_row_map = {}
    fr = 4

    def fc_section(title):
        nonlocal fr
        for col_i in range(fc_years + 1):
            c = ws_f.cell(row=fr, column=col_i + 1, value=title if col_i == 0 else None)
            c.font = Font(name="Arial", bold=True, color=WHITE_TEXT, size=10)
            c.fill = _fill(HEADER_BG)
        fr += 1

    def fc_row(label, key, formula_fn, fmt=FMT_CURRENCY):
        nonlocal fr
        ws_f.cell(row=fr, column=1, value=label).font = _font()
        fc_row_map[key] = fr
        for i in range(fc_years):
            cl = get_column_letter(i + 2)
            prev_cl = get_column_letter(i + 1)
            f = formula_fn(i, cl, prev_cl)
            c = ws_f.cell(row=fr, column=i + 2, value=f)
            c.font = _font(color=BLACK_FORM)
            c.number_format = fmt
        fr += 1

    def fc_blank():
        nonlocal fr
        fr += 1

    fc_section("── INCOME STATEMENT ──")

    # Revenue
    last_rev_row = rmap.get("revenue")
    last_col     = get_column_letter(n + 1)  # last historical column

    def rev_formula(i, cl, prev_cl):
        if i == 0:
            return f"=Historical_Model!{last_col}{last_rev_row}*(1+{cl}{rev_asmp_r})"
        return f"={prev_cl}{fc_row_map['rev']}*(1+{cl}{rev_asmp_r})"

    fc_row("Revenue", "rev", rev_formula)
    rev_fc_r = fc_row_map["rev"]

    fc_row("Cost of Goods Sold", "cogs_fc", lambda i, cl, _: f"={cl}{rev_fc_r}*{cl}{cogs_asmp_r}")
    cogs_fc_r = fc_row_map["cogs_fc"]

    fc_row("Gross Profit", "gp_fc", lambda i, cl, _: f"={cl}{rev_fc_r}-{cl}{cogs_fc_r}")
    gp_fc_r = fc_row_map["gp_fc"]

    fc_row("Gross Margin %", "gm_fc", lambda i, cl, _: f"={cl}{gp_fc_r}/{cl}{rev_fc_r}", fmt=FMT_PCT)

    fc_blank()
    fc_row("Total OPEX", "opex_fc", lambda i, cl, _: f"={cl}{rev_fc_r}*{cl}{opex_asmp_r}")
    opex_fc_r = fc_row_map["opex_fc"]

    fc_row("OPEX % of Revenue", "opex_pct_fc", lambda i, cl, _: f"={cl}{opex_fc_r}/{cl}{rev_fc_r}", fmt=FMT_PCT)

    fc_blank()
    fc_row("Operating Income (EBIT)", "ebit_fc", lambda i, cl, _: f"={cl}{gp_fc_r}-{cl}{opex_fc_r}")
    ebit_fc_r = fc_row_map["ebit_fc"]

    fc_row("EBIT Margin %", "ebit_pct_fc", lambda i, cl, _: f"={cl}{ebit_fc_r}/{cl}{rev_fc_r}", fmt=FMT_PCT)

    fc_blank()
    fc_section("── CASH FLOW ──")

    fc_row("Operating Cash Flow (est.)", "ocf_fc", lambda i, cl, _: f"={cl}{ebit_fc_r}*0.85")
    ocf_fc_r = fc_row_map["ocf_fc"]

    fc_row("Capital Expenditures", "capex_fc", lambda i, cl, _: f"={cl}{rev_fc_r}*{cl}{cap_asmp_r}")
    capex_fc_r = fc_row_map["capex_fc"]

    fc_row("Free Cash Flow", "fcf_fc", lambda i, cl, _: f"={cl}{ocf_fc_r}-{cl}{capex_fc_r}")

    # ── Scenarios ────────────────────────────────────────────────────────────
    ws_sc = wb.create_sheet("Scenarios")
    _set_col_widths(ws_sc, {"A": 35, "B": 18, "C": 18, "D": 18, "E": 40})
    ws_sc.cell(row=1, column=1, value=f"{company} — Scenario Analysis (Year 1 Forecast)").font = _font(bold=True, size=12)
    _header_row(ws_sc, 2, ["Assumption", "Downside", "Base Case", "Upside", "Notes"])

    sc_asmp = [
        ("Revenue Growth Rate",   -0.10,          0.05,          0.15,     "Base=5%; Down=-10%; Up=+15%"),
        ("COGS % of Revenue",     round(defaults[1]*1.05, 3), round(defaults[1], 3), round(defaults[1]*0.95, 3), "±5% of Base"),
        ("OPEX % of Revenue",     round(defaults[2]*1.10, 3), round(defaults[2], 3), round(defaults[2]*0.90, 3), "±10% of Base"),
        ("Capex % of Revenue",    round(defaults[3]*1.20, 3), round(defaults[3], 3), round(defaults[3]*0.80, 3), "±20% of Base"),
    ]

    for i, (lbl, down, base, up, note) in enumerate(sc_asmp, start=3):
        ws_sc.cell(row=i, column=1, value=lbl).font = _font()
        for col_i, val in enumerate([down, base, up], start=2):
            c = ws_sc.cell(row=i, column=col_i, value=val)
            c.font = _font(color=BLUE_INPUT)
            c.number_format = FMT_PCT
        ws_sc.cell(row=i, column=5, value=note).font = _font()

    # Output rows
    ws_sc.cell(row=8, column=1, value="── PROJECTED OUTPUTS (Year 1) ──").font = Font(name="Arial", bold=True, color=WHITE_TEXT)
    ws_sc.cell(row=8, column=1).fill = _fill(HEADER_BG)
    for col_i in range(2, 6):
        ws_sc.cell(row=8, column=col_i).fill = _fill(HEADER_BG)

    base_rev = rev_last
    sc_outputs = [
        ("Revenue", [base_rev*(1+sc_asmp[0][1]), base_rev*(1+sc_asmp[0][2]), base_rev*(1+sc_asmp[0][3])]),
        ("Gross Profit", [
            base_rev*(1+sc_asmp[0][1])*(1-sc_asmp[1][1]),
            base_rev*(1+sc_asmp[0][2])*(1-sc_asmp[1][2]),
            base_rev*(1+sc_asmp[0][3])*(1-sc_asmp[1][3])]),
        ("EBIT", [
            base_rev*(1+sc_asmp[0][1])*((1-sc_asmp[1][1])-sc_asmp[2][1]),
            base_rev*(1+sc_asmp[0][2])*((1-sc_asmp[1][2])-sc_asmp[2][2]),
            base_rev*(1+sc_asmp[0][3])*((1-sc_asmp[1][3])-sc_asmp[2][3])]),
    ]

    for i, (lbl, vals) in enumerate(sc_outputs, start=9):
        ws_sc.cell(row=i, column=1, value=lbl).font = _font()
        for col_i, val in enumerate(vals, start=2):
            c = ws_sc.cell(row=i, column=col_i, value=val)
            c.font = _font(color=BLACK_FORM)
            c.number_format = FMT_CURRENCY

    # ── Validation ───────────────────────────────────────────────────────────
    ws_v = wb.create_sheet("Validation")
    _set_col_widths(ws_v, {"A": 45, "B": 20, "C": 20, "D": 15})
    ws_v.cell(row=1, column=1, value="Validation Checks").font = _font(bold=True, size=12)
    _header_row(ws_v, 2, ["Check", "Computed", "Expected (Source)", "Status"])

    checks = []
    for d in periods_data:
        per = d.get("period", "?")
        rev = _safe(d.get("revenue"))
        cgs = _safe(d.get("cogs"))
        gp  = _safe(d.get("gross_profit"))
        totopex = _safe(d.get("total_opex"))
        ebit = _safe(d.get("operating_income"))
        ni   = _safe(d.get("net_income"))
        ocf  = _safe(d.get("ocf"))
        cap  = _safe(d.get("capex"))

        if rev and cgs:
            checks.append((f"{per}: Gross Profit = Revenue - COGS", rev - cgs, gp))
        if gp and totopex:
            checks.append((f"{per}: EBIT = Gross Profit - Total OPEX", gp - totopex, ebit))
        if ocf is not None and cap:
            checks.append((f"{per}: FCF = OCF - Capex", ocf - cap, None))

    for i, (lbl, computed, expected) in enumerate(checks, start=3):
        ws_v.cell(row=i, column=1, value=lbl).font = _font()
        ws_v.cell(row=i, column=2, value=computed).number_format = FMT_CURRENCY
        ws_v.cell(row=i, column=2).font = _font(color=BLACK_FORM)
        if expected is not None:
            ws_v.cell(row=i, column=3, value=expected).number_format = FMT_CURRENCY
            ws_v.cell(row=i, column=3).font = _font(color=BLUE_INPUT)
            diff = abs(computed - expected) if computed is not None else 999
            status = "✓ PASS" if diff < 10 else f"⚠ DIFF {diff:,.0f}"
            c = ws_v.cell(row=i, column=4, value=status)
            c.font = Font(name="Arial", bold=True, color="FF008000" if "PASS" in status else "FFFF0000")
        else:
            ws_v.cell(row=i, column=4, value="ℹ Computed").font = _font()

    # ── Commentary ───────────────────────────────────────────────────────────
    ws_c = wb.create_sheet("Commentary")
    _set_col_widths(ws_c, {"A": 20, "B": 100})
    ws_c.cell(row=1, column=1, value="Model Commentary").font = _font(bold=True, size=12)
    _header_row(ws_c, 2, ["Period", "Observation"])

    for i, d in enumerate(periods_data, start=3):
        per = d.get("period", "")
        rev = _safe(d.get("revenue"))
        ni  = _safe(d.get("net_income"))
        gm  = d.get("gross_margin_pct") or ((_safe(d.get("gross_profit")) / rev) if rev else None)
        note = d.get("notes", "")
        obs = f"Revenue: {d['currency'] if d.get('currency') else '$'}{rev:,.0f}  |  Net Income: {ni:,.0f}  |  Gross Margin: {gm:.1%}" if gm else f"Revenue: {rev:,.0f}  |  Net Income: {ni:,.0f}"
        if note:
            obs += f"  |  Notes: {note}"
        ws_c.cell(row=i, column=1, value=per).font = _font(bold=True)
        c = ws_c.cell(row=i, column=2, value=obs)
        c.font = _font()
        c.alignment = Alignment(wrap_text=True)

    ws_c.cell(row=len(periods_data)+3, column=1, value="Forecast Logic").font = _font(bold=True)
    ws_c.cell(row=len(periods_data)+3, column=2, value="Revenue grows from the last historical period. COGS% and OPEX% default to historical averages. Edit assumption cells on the Forecast sheet (rows 29-32) to update all forecast outputs dynamically.").font = _font()

    # ── Sources ──────────────────────────────────────────────────────────────
    ws_src = wb.create_sheet("Sources")
    _set_col_widths(ws_src, {"A": 15, "B": 60, "C": 20, "D": 20})
    ws_src.cell(row=1, column=1, value="Sources").font = _font(bold=True, size=12)
    _header_row(ws_src, 2, ["Period", "Source File", "Audit Status", "Extraction Date"])

    for i, d in enumerate(periods_data, start=3):
        ws_src.cell(row=i, column=1, value=d.get("period", "")).font = _font()
        ws_src.cell(row=i, column=2, value=d.get("source_file", "")).font = _font()
        ws_src.cell(row=i, column=3, value=d.get("audit_status", "")).font = _font()
        ws_src.cell(row=i, column=4, value=datetime.now().strftime("%Y-%m-%d")).font = _font()

    # ── Raw_Extract (combined) ───────────────────────────────────────────────
    ws_r = wb.create_sheet("Raw_Extract")
    raw_headers = ["source_file", "period", "statement_type", "original_label",
                   "normalized_label", "value", "unit", "confidence"]
    _header_row(ws_r, 1, raw_headers)
    _set_col_widths(ws_r, {"A": 50, "D": 40, "E": 30, "F": 18})

    raw_r = 2
    for d in periods_data:
        sf   = d.get("source_file", "")
        per  = d.get("period", "")
        unit = d.get("unit", "$")
        items = [
            ("Income Statement", "Revenue", "Revenue", d.get("revenue")),
            ("Income Statement", "Cost of goods sold", "COGS", d.get("cogs")),
            ("Income Statement", "Gross profit", "Gross Profit", d.get("gross_profit")),
            ("Income Statement", "Operating expenses", "Operating Expenses", d.get("opex_operating")),
            ("Income Statement", "D&A", "D&A", d.get("opex_amortization")),
            ("Income Statement", "Total OPEX", "Total OPEX", d.get("total_opex")),
            ("Income Statement", "Operating income / EBIT", "EBIT", d.get("operating_income")),
            ("Income Statement", "Government assistance", "Govt Assistance", d.get("government_assistance")),
            ("Income Statement", "Net income", "Net Income", d.get("net_income")),
            ("Cash Flow", "Operating cash flow", "OCF", d.get("ocf")),
            ("Cash Flow", "Capital expenditures", "Capex", d.get("capex")),
            ("Cash Flow", "Investing cash flow", "Investing CF", d.get("investing_cf")),
            ("Cash Flow", "Financing cash flow", "Financing CF", d.get("financing_cf")),
            ("Balance Sheet", "Cash", "Cash", d.get("cash")),
            ("Balance Sheet", "Accounts receivable", "AR", d.get("accounts_receivable")),
            ("Balance Sheet", "Inventories", "Inventory", d.get("inventories")),
            ("Balance Sheet", "PP&E net", "PP&E", d.get("capital_assets")),
            ("Balance Sheet", "Total assets", "Total Assets", d.get("total_assets")),
            ("Balance Sheet", "Accounts payable", "AP", d.get("accounts_payable")),
            ("Balance Sheet", "Total debt", "Total Debt", d.get("total_debt")),
            ("Balance Sheet", "Total equity", "Total Equity", d.get("total_equity")),
        ]
        for stmt, orig, norm, val in items:
            row_data = [sf, per, stmt, orig, norm, val, unit, "High"]
            for col_i, cell_val in enumerate(row_data, start=1):
                ws_r.cell(row=raw_r, column=col_i, value=cell_val).font = _font()
            ws_r.cell(row=raw_r, column=6).number_format = FMT_CURRENCY
            raw_r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
